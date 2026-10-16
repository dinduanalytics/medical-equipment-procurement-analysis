"""Create a public-safe copy of the procurement workbook.

Usage:
    python scripts/sanitize_dataset.py "Procurement Dataset.xlsx"

The script reads the 'Procurement Data' worksheet and writes only the cleaned
procurement table. Supplier contact names, phone numbers, email addresses,
blank columns, and other worksheets are excluded from the public output.
"""
from pathlib import Path
import sys
from openpyxl import load_workbook, Workbook

SENSITIVE_HEADERS = {
    "supplier contact",
    "supplier phone",
    "supplier email",
}
SOURCE_SHEET = "Procurement Data"
OUTPUT_PATH = Path("data/procurement_dataset_sanitized.xlsx")


def clean_header(value):
    return " ".join(str(value or "").split()).strip()


def main():
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("Procurement Dataset.xlsx")
    if not source.exists():
        raise FileNotFoundError(
            f"Source workbook not found: {source}. Pass its path as the first argument."
        )

    source_wb = load_workbook(source, data_only=True, read_only=True)
    if SOURCE_SHEET not in source_wb.sheetnames:
        raise ValueError(f"Expected worksheet '{SOURCE_SHEET}' was not found.")

    source_ws = source_wb[SOURCE_SHEET]
    raw_headers = next(source_ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [clean_header(h) for h in raw_headers]
    keep_indexes = [
        i for i, header in enumerate(headers)
        if header and header.casefold() not in SENSITIVE_HEADERS
    ]

    output_wb = Workbook()
    output_ws = output_wb.active
    output_ws.title = "Procurement Data"
    output_ws.append([headers[i] for i in keep_indexes])

    for row in source_ws.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in row):
            continue
        output_ws.append([row[i] if i < len(row) else None for i in keep_indexes])

    output_ws.freeze_panes = "A2"
    output_ws.auto_filter.ref = output_ws.dimensions
    for column_cells in output_ws.columns:
        max_length = max((len(str(cell.value)) if cell.value is not None else 0)
                         for cell in column_cells)
        output_ws.column_dimensions[column_cells[0].column_letter].width = min(max_length + 2, 34)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(OUTPUT_PATH)
    print(f"Sanitized workbook saved to: {OUTPUT_PATH}")
    print(f"Rows written (excluding header): {output_ws.max_row - 1}")
    print(f"Columns retained: {output_ws.max_column}")
    print("Removed fields: Supplier Contact, Supplier Phone, Supplier Email")


if __name__ == "__main__":
    main()
